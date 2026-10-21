"""Expand the published 100-question bank without changing existing stable IDs."""
import json, csv, hashlib
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo
root=Path(__file__).resolve().parents[1]
bank=json.loads((root/'public/data/halloween-100-v2.json').read_text(encoding='utf-8'))
items=bank['questions'][:100]
# Each authored entry: prompt, correct choice, three distractors, explanation.
authored='''
我是派對的紙裝飾：有八隻腳，常被掛在假蜘蛛網上。猜猜我是誰？|蜘蛛|蝙蝠|蝴蝶|蝸牛|蜘蛛有八隻腳；蝙蝠、蝴蝶與蝸牛都不符合八隻腳的線索。
白天收起翅膀，夜裡飛出來；我是哺乳類，卻不是鳥。猜猜我是誰？|蝙蝠|貓頭鷹|烏鴉|麻雀|蝙蝠是會飛行的哺乳類，另外三個選項都是鳥類。
巫師把會融化的派對裝飾放在暖爐旁。哪一個最容易變成水？|冰雕南瓜|紙南瓜|陶瓷南瓜|木南瓜|冰受熱會融化成水，紙、陶瓷與木頭不會像冰一樣融化成水。
南瓜餅乾要裝袋送人，哪個步驟能保留食材資訊？|在袋上標示成分|只畫一隻蝙蝠|只寫袋子顏色|把標籤藏起來|成分標示能讓收禮者知道餅乾用了什麼，圖案或顏色無法取代成分資訊。
幽靈的紙帽被風吹跑了。風主要是在移動什麼？|空氣|影子|月亮|聲音的顏色|風是空氣的流動；影子、月亮和不存在的聲音顏色都不是風本身。
哪個萬聖節裝飾不需要通電，也不必點火就能被風吹動？|懸掛的紙蝙蝠|插電的投影機|電動南瓜車|LED燈串|輕薄且懸掛的紙蝙蝠能被風吹動，其餘選項需要電力才能執行原本功能。
吸血鬼的鏡子結了一層水珠，最適合用什麼擦乾？|乾淨軟布|一把沙子|金屬叉子|糖漿|軟布能吸水並減少刮傷；沙子與叉子可能刮傷，糖漿會讓鏡面更黏。
鬼屋用手電筒照紙剪影。哪種材料最能讓光直接穿過？|透明塑膠片|厚紙板|木板|金屬片|透明塑膠片能透光，厚紙板、木板和金屬片會擋住大部分光。
哪一樣是南瓜原本生長時會有的部分？|種子|拉鍊|電池盒|玻璃眼珠|南瓜果實裡有種子，拉鍊、電池盒與玻璃眼珠都是人為添加物。
蝙蝠披風弄濕了，想晾乾應放哪裡？|通風的晾衣架|密封塑膠盒|裝滿水的盆子|冰箱冷凍層|通風有助水分蒸發；密封、泡水或冷凍都不是合適的晾乾方法。
派對邀請函寫著「請於入口排隊」。應怎麼做？|依序等前面的人進入|從隊伍中間插進去|堵住入口聊天|逆著人群衝進去|依序等候符合排隊要求，也能維持入口通暢。
南瓜拼圖少一塊，缺口是三角形，應找哪塊？|三角形拼片|圓形拼片|正方形拼片|五角星拼片|拼片外形必須吻合缺口；三角形缺口應配三角形拼片。
幽靈主持人說「先洗手，再拿點心」。第一步是什麼？|洗手|拿點心|拆糖紙|吃餅乾|題目明確把洗手排在拿點心之前，其他動作都不是第一步。
舞台突然停電，哪個東西能自己發出可見光？|已開啟的電池手電筒|未貼反光片的黑布|普通木湯匙|空紙箱|有電且已開啟的手電筒會發光，黑布、木湯匙和紙箱不會自行發光。
糖果屋的地圖想標示「不可進入」，哪個圖示最清楚？|入口加上禁止符號|入口畫一顆愛心|入口畫一朵花|入口寫歡迎|禁止符號能表達不能進入；愛心、花與歡迎不會傳達禁止的意思。
裝扮成機器人的朋友說話太小聲，怎樣比較容易溝通？|走近並請他再說一次|大家一起大叫|直接猜他說什麼|背對他跑開|走近並確認內容有助聽清楚；大叫、猜測或離開都無法確認原話。
派對桌上灑了糖漿，哪個訊號最適合提醒大家？|放上小心黏滑的提示牌|撒更多糖漿|用布遮住不說|把燈全部關掉|清楚的提示能提醒路過者注意，增加糖漿、遮住或關燈會使問題更難察覺。
黑貓造型氣球越吹越大，裡面增加了什麼？|空氣|黑色顏料|貓毛|影子|吹氣球會往裡面加入空氣，外觀變大不是因為顏料、貓毛或影子增加。
南瓜湯與南瓜餅都要用到哪種主要食材？|南瓜|海苔|可可豆|檸檬皮|兩個名稱都指向南瓜；海苔、可可豆和檸檬皮不是題目指定的共同主要食材。
怪獸手套左右各一隻，一組手套共有幾隻？|兩隻|一隻|三隻|四隻|左手一隻加右手一隻，合計兩隻，不能把手指數當成手套數。
萬聖節英文 Halloween 裡，哪個字母出現兩次且連在一起？|l|H|a|n|Halloween 中間有連續的 ll；H、a、n 沒有連續出現兩次。
紙幽靈的影子需要什麼條件才能出現？|有光照向被遮擋的表面|房間完全沒有光|只有一段錄音|只有一個氣味|影子來自物體擋住光；完全無光、錄音或氣味都不能單獨產生可見影子。
巫師想量南瓜的重量，該拿哪個工具？|秤|尺|溫度計|計時器|秤測量重量；尺量長度、溫度計量溫度、計時器量時間。
要量蝙蝠披風的長度，哪個工具最合適？|捲尺|體重計|量杯|指南針|捲尺量長度；體重計、量杯與指南針分別用於重量、容量與方向。
想知道南瓜餅烤了多久，應看哪個工具？|計時器|放大鏡|秤|量角器|計時器記錄經過的時間；放大鏡、秤與量角器不測量烘烤時間。
哪個裝扮最符合「全身畫著黑白條紋的馬」？|斑馬|長頸鹿|企鵝|瓢蟲|斑馬是有黑白條紋的馬科動物，其他選項不符合全部線索。
貓頭鷹面具的兩個眼孔應該留在哪裡？|配戴者眼睛對應的位置|帽子頂端|鞋底|背後|眼孔必須對準眼睛才能看路，帽頂、鞋底和背後都無法提供正常視線。
派對用可重複使用的杯子，結束後怎麼處理最合理？|清洗晾乾後收好|每個都打破|全部丟進河裡|藏在點心裡|可重複使用的杯子清洗後能繼續使用，破壞、亂丟或藏入食物都不適合。
南瓜燈劇場裡，誰負責按照劇本演出角色？|演員|觀眾座椅|售票信封|垃圾桶|演員負責扮演角色；座椅、信封與垃圾桶都不是演出者。
哪一種紙裝飾折好後最像「有兩片翅膀的小動物」？|紙蝙蝠|紙南瓜|紙帽子|紙糖果|蝙蝠有兩片翅膀；南瓜、帽子與糖果沒有動物翅膀。
黑貓踩過顏料，地板出現一串什麼最合理？|腳印|車輪印|魚鱗|樹葉脈|貓用腳走路，沾顏料後會留下腳印，不會留下車輪印、魚鱗或葉脈。
巫師把兩杯不同顏色的彩珠混在一起，哪種方法能按顏色分開？|依顏色逐顆分類|把眼睛閉起來亂抓|只數杯子不看珠子|再加入更多顏色|按顏色觀察並分類能分開彩珠，亂抓、數杯子或增加顏色不能完成指定分類。
幽靈紙偶需要一個圓形頭部，哪個物品的外框可用來描圓？|圓杯口|三角尺外框|方盒底部|星形模具|圓杯口的外框是圓形，其他工具描出的外框分別是三角形、方形與星形。
哪一個派對告示是在邀請大家安靜？|請降低音量|請大聲尖叫|一起敲鼓|歡迎吹哨|降低音量表示聲音變小，尖叫、敲鼓和吹哨都會增加聲音。
要把小紙蝙蝠固定在卡片上，哪個工具最適合？|膠水|量杯|手電筒|指南針|膠水能黏住紙張，量杯、手電筒與指南針不能取代黏著功能。
哪種南瓜臉只有三條邊的嘴巴？|三角形嘴巴|圓形嘴巴|正方形嘴巴|六邊形嘴巴|三角形有三條邊；圓形沒有直邊，正方形四邊，六邊形六邊。
兩個人背對背站著，一人看前方，另一人看哪裡？|相反方向|完全相同方向|一定看天空|一定看地面|背對背時朝向相反；不代表必須抬頭或低頭。
舞會面具上的亮片反射燈光，關掉所有光源後會怎樣？|一般亮片不再反射出可見亮光|一定變成手電筒|一定自己唱歌|立刻變成糖果|普通亮片需要外來光才能反射出亮光，不會因關燈產生電力或變成別的物品。
邀請函說「請帶一樣橘色物品」。哪個符合？|橘色紙南瓜|紫色披風|白色面具|綠色手套|題目只要求橘色，只有橘色紙南瓜符合指定顏色。
幽靈想找回原本的南瓜圖案，應先比對哪項特徵？|形狀與花紋|附近人的名字|今天星期幾|房間外的風向|形狀與花紋是圖案本身的特徵，名字、星期與風向不能直接辨認圖案。
'''.strip().splitlines()
def add(category,prompt,correct,wrongs,explanation):
 ident=f'halloween-extra-{len(items)+1:03d}'
 # Rotate exported answer positions; gameplay still independently shuffles IDs.
 texts=[correct,*wrongs]; offset=len(items)%4;texts=texts[offset:]+texts[:offset]
 options=[{'id':f'{ident}.option-{i+1}','text':t,**({'wrongReason':f'「{t}」不符合條件。{explanation}'} if t!=correct else {})} for i,t in enumerate(texts)]
 items.append({'id':ident,'category':category,'prompt':prompt,'options':options,'answer':{'id':ident+'.answer','optionId':next(o['id'] for o in options if o['text']==correct)},'hint':{'id':ident+'.hint','text':explanation}})
for line in authored:
 prompt,correct,a,b,c,explanation=line.split('|');add('萬聖節觀察與小知識',prompt,correct,[a,b,c],explanation)
assert len(items)==140,len(items)
# 160 independently solvable themed puzzles across eight forms.
for n in range(1,21):
 a=n+3;b=n%5+2
 def number(cat,prompt,answer,wrong,why,unit=''):
  add(cat,prompt,str(answer)+unit,[str(x)+unit for x in wrong],why)
 number('糖果補給',f'第{n}桌原有{a}顆南瓜糖，又補上{b}顆，最後送出2顆。桌上剩幾顆？',a+b-2,[a+b,a+b-1,a+b+1],f'先加上補來的，再扣掉送出的：{a}＋{b}－2＝{a+b-2}。','顆')
 number('怪獸分糖',f'{b}個怪獸糖果袋各放{a}顆糖。把所有袋子倒出來，共有幾顆？',a*b,[a*b-1,a*b+1,a*b+b],f'{b}袋各{a}顆，合計{a}×{b}＝{a*b}顆。','顆')
 total=(n+2)*3
 number('派對平均分',f'南瓜托盤上有{total}塊餅乾，要平均裝進3個幽靈點心盒，每盒幾塊？',n+2,[n+1,n+3,total],f'平均分成3盒，每盒{total}÷3＝{n+2}塊。','塊')
 step=n%4+2;start=n+1;seq=[start+i*step for i in range(3)];ans=start+3*step
 number('南瓜燈數字規律',f'第{n}排燈籠編號依序為{seq[0]}、{seq[1]}、{seq[2]}，每次增加同樣的數。下一號是多少？',ans,[ans-1,ans+1,ans+step],f'相鄰兩號都增加{step}，所以{seq[2]}＋{step}＝{ans}。')
 # Finite one-based cyclic sequences, varied motifs and positions.
 motifs=[['南瓜','蝙蝠','星星'],['帽子','掃帚','月亮','黑貓'],['幽靈','糖果','蜘蛛','燈籠','面具']][n%3]
 count=len(motifs);pos=n+count+2;correct=motifs[(pos-1)%count]
 wrong=[x for x in ['南瓜','蝙蝠','星星','帽子','掃帚','月亮','黑貓','幽靈','糖果','蜘蛛','燈籠','面具'] if x!=correct][:3]
 add('派對圖案循環',f'第{n}串掛飾從第1個開始，按照「'+ '→'.join(motifs)+f'」一直重複。第{pos}個圖案是？',correct,wrong,f'每{count}個一組，第{pos}個和組內第{(pos-1)%count+1}個相同，是{correct}。')
 number('面具配對',f'第{n}場舞會來了{a}位客人，每兩人組成一隊；若有人落單就先休息。最多能組幾隊？',a//2,[a//2+1,a,a+1],f'每隊要兩人，{a}÷2可以組{a//2}個完整隊伍'+('，另外1人先休息。' if a%2 else '，剛好沒有剩下的人。'),'隊')
 # Position logic: unambiguous fixed order and reverse viewpoint.
 length=n+5;position=n%4+2;ans=length-position+1
 number('遊行換個方向數',f'{length}盞南瓜燈排成一直線，一盞紫帽燈從左邊數是第{position}盞。從右邊數是第幾盞？',ans,[ans-1,ans+1,ans+2],f'總數減去左邊順位再加1：{length}－{position}＋1＝{ans}。')
 # All units explicit, no date or cultural accuracy dependence.
 time=n%6+2;dur=n%4+3;ans=time+dur
 number('幽靈舞台計時',f'第{n}段表演從計時器顯示第{time}分鐘時開始，演出{dur}分鐘後結束。結束時顯示第幾分鐘？',ans,[ans-1,ans+1,ans+2],f'開始時間加上經過時間：{time}＋{dur}＝{ans}分鐘。')
assert len(items)==300
assert len({q['prompt'] for q in items})==300
bank.update(version='halloween-300-v1',title='小貝爾與火馬・萬聖節300題',questions=items)
bank['contentSha256']=hashlib.sha256(json.dumps(items,ensure_ascii=False,separators=(',',':')).encode()).hexdigest()
for path in [root/'public/data/halloween-300-v1.json',root/'cloud-question-bank/halloween-300-v1.json']:
 path.write_text(json.dumps(bank,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
headers=['啟用','題目ID（勿改）','分類','題目','選項A','選項B','選項C','選項D','正確答案','提示','A錯因','B錯因','C錯因','D錯因','A_ID（勿改）','B_ID（勿改）','C_ID（勿改）','D_ID（勿改）','標準答案ID（勿改）','提示ID（勿改）']
rows=[headers]
for q in items:
 opts=q['options'];rows.append(['是',q['id'],q.get('category',''),q['prompt'],*[o['text'] for o in opts],'ABCD'[[o['id'] for o in opts].index(q['answer']['optionId'])],q['hint']['text'],*[o.get('wrongReason','') for o in opts],*[o['id'] for o in opts],q['answer']['id'],q['hint']['id']])
with (root/'cloud-question-bank/萬聖節300題-維護表.csv').open('w',encoding='utf-8-sig',newline='') as f:csv.writer(f).writerows(rows)
(root/'cloud-question-bank/maintenance-300.tsv').write_text('\n'.join('\t'.join(str(v).replace('\n',' ').replace('\t',' ') for v in r) for r in rows),encoding='utf-8')
(root/'cloud-question-bank/append-200.tsv').write_text('\n'.join('\t'.join(str(v).replace('\n',' ').replace('\t',' ') for v in r) for r in rows[101:]),encoding='utf-8')
wb=Workbook();ws=wb.active;ws.title='題庫'
for row in rows:ws.append(row)
ws.freeze_panes='E2';ws.auto_filter.ref=ws.dimensions
for cell in ws[1]:cell.fill=PatternFill('solid',fgColor='38204D');cell.font=Font(color='FFFFFF',bold=True)
for row in ws.iter_rows(min_row=2):
 for cell in row:cell.alignment=Alignment(vertical='top',wrap_text=True)
for col in ['B','O','P','Q','R','S','T']:ws.column_dimensions[col].hidden=True
for col in ['D','J','K','L','M','N']:ws.column_dimensions[col].width=48
for col in ['E','F','G','H']:ws.column_dimensions[col].width=24
ws.column_dimensions['C'].width=23
for i in range(2,302):ws.row_dimensions[i].height=66
notes=wb.create_sheet('使用說明')
for row in [('項目','說明'),('題數','300題，保留原100題穩定ID，新增200題。'),('每關','抽取3題不重複；選項位置另外洗牌。'),('正確答案','以A/B/C/D表示此表格的原始選項位置。'),('錯因','三個錯誤選項都必須填原因。'),('穩定ID','隱藏欄位勿改；更改題幹、選項、提示不必改引擎。'),('發布','雲端或Excel編輯不會自動發布；需匯出資料、更新正式JSON、重新建置與推送。'),('版本','halloween-300-v1')]:notes.append(row)
notes.column_dimensions['A'].width=20;notes.column_dimensions['B'].width=100
for row in notes:
 for cell in row:cell.alignment=Alignment(wrap_text=True,vertical='top')
wb.save(root/'cloud-question-bank/萬聖節300題-維護主表.xlsx')
print('300 questions; 1200 options; 900 wrong reasons; original 100 IDs unchanged')
